"""Construye el reporte ejecutivo de la Entrega 1 (docx, Arial 12) a partir de los resultados del EDA.

Las cifras se leen de reports/tables/resultados_eda.json y linea_base_cv.csv para garantizar
consistencia entre el notebook y el reporte. Uso: python src/build_report.py
"""
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
R = json.load(open(ROOT / "reports/tables/resultados_eda.json", encoding="utf-8"))
LB = pd.read_csv(ROOT / "reports/tables/linea_base_cv.csv")
REPO = "https://github.com/Juliethr3/neuroagenda-minhalma"

def pct(x, d=1):
    return f"{x * 100:.{d}f} %".replace(".", ",")

def num(x, d=2):
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")

def ent(x):
    return f"{x:,}".replace(",", ".")

OR = R["or_ajustados"]
def orf(k):
    o = OR[k]
    return f"OR = {num(o['OR'])} [IC 95 %: {num(o['ic95_inf'])}–{num(o['ic95_sup'])}]"

best = LB[(LB.variables == "Con historial del paciente") & (LB.modelo.str.contains("HGB"))].iloc[0]
lr = LB[(LB.variables == "Sin historial") & (LB.modelo.str.contains("logística"))].iloc[0]
ta = R["tasa_por_antelacion"]; te = R["tasa_por_edad"]; th = R["tasa_por_inasist_previas"]

# ----------------------------------------------------------------------------- estilos
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(27.94), Cm(21.59)  # carta
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.0))
sec.top_margin = sec.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(3); st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0
for s in ("List Bullet",):
    doc.styles[s].font.name = "Arial"; doc.styles[s].font.size = Pt(12)

def para(text="", bold_lead=None, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False, after=3):
    p = doc.add_paragraph(); p.alignment = align; p.paragraph_format.space_after = Pt(after)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True; r.font.size = Pt(size)
    _runs(p, text, size, italic)
    return p

def _runs(p, text, size, italic=False):
    # soporta **negrita** y [[resaltado]] (marcadores a completar por el equipo)
    import re
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\[\[.+?\]\])", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True; r.font.size = Pt(size); continue
        elif tok.startswith("[["):
            r = p.add_run(tok[2:-2]); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = p.add_run(tok)
        r.font.size = Pt(size); r.italic = italic

def bullet(text, size=12):
    p = doc.add_paragraph(style="List Bullet"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.left_indent = Cm(0.6)
    _runs(p, text, size)

def h(text, lvl=1):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6 if lvl == 1 else 3)
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F) if lvl == 1 else RGBColor(0x33, 0x33, 0x33)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

def table(rows, widths, size=10, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            _runs(p, str(val), size)
            if header and i == 0:
                for r in p.runs: r.bold = True
                shade(c, "DCE6F1")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def figure(path, width_cm, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    para(caption, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, italic=True, after=4)

# ----------------------------------------------------------------------------- encabezado
p = para("", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
r = p.add_run("NeuroAgenda: predicción del riesgo de inasistencia en consulta ambulatoria de psicología, nutrición y fisioterapia")
r.bold = True; r.font.size = Pt(14)
para("MINE-4101 Ciencia de Datos Aplicada · Proyecto final · Entrega 1 · Octubre 2026 · Equipo: [[Nombre integrante 1; Nombre integrante 2]] · "
     f"Repositorio: {REPO} · Video: [[enlace YouTube]]", size=10, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)

# ----------------------------------------------------------------------------- 1
h("1. Problemática y entendimiento del negocio")
para("**Organización.** **MinHalma**, centro de salud ambulatorio de reciente creación que ofrece psicología, nutrición y fisioterapia. "
     "**Stakeholder:** [[nombre y cargo]]. Por ser un centro nuevo, su historial de citas aún es corto: esto justifica validar el enfoque con un "
     "conjunto público (sección 5) y es, a la vez, una oportunidad para estandarizar desde el inicio el registro del estado de cada cita.")
para("**Problema.** Una de cada cinco citas ambulatorias no se cumple: la tasa promedio de inasistencia (no-show) reportada en 105 estudios es "
     "cercana al 23 % (Dantas et al., 2018). En Colombia, el Hospital Regional de San Gil reportó 15.464 citas incumplidas solo en el primer "
     "semestre de 2025 (Blu Radio, 2025), y la ECV 2025 del DANE muestra que la proporción de personas con un problema de salud que acudieron a su "
     "entidad bajó de 63,3 % a 56,8 % entre 2024 y 2025 (Valora Analitik, 2026). Psicología, nutrición y fisioterapia funcionan por sesiones "
     "recurrentes: cada inasistencia es tiempo profesional no facturado, un cupo que otro paciente no pudo usar y una ruptura de la continuidad "
     "del tratamiento. Para un centro nuevo, la ocupación de agenda es además una condición de sostenibilidad financiera. Hoy los recordatorios son indiscriminados: se contacta igual a quien siempre asiste que a quien probablemente faltará.")
para("**Oportunidad y objetivo.** Estimar, al momento de agendar y 72 h antes de la cita, la probabilidad individual de inasistencia para "
     "**focalizar las acciones de confirmación y reasignación de cupos** en las citas de mayor riesgo. Objetivos específicos: (O1) construir y "
     "validar un modelo calibrado de riesgo; (O2) exponerlo mediante una API REST; (O3) entregar un tablero operativo con la lista priorizada y los KPI.")
table([
    ["KPI de negocio", "Definición", "Meta a 6 meses (propuesta)"],
    ["Tasa de inasistencia", "Citas no asistidas / citas programadas", "Reducción relativa ≥ 20 % frente a la línea base"],
    ["Cupos recuperados", "Cupos liberados con antelación y reasignados", "≥ 50 % de las cancelaciones anticipadas"],
    ["Ocupación efectiva", "Horas de especialista atendidas / horas ofertadas", "+5 p.p."],
    ["Eficiencia de contacto", "Inasistencias evitadas por cada 100 contactos", "Superior a la estrategia de contactar al azar"],
    ["Equidad", "Brecha de recall del modelo entre grupos de edad, sexo y subsidio", "≤ 5 p.p."],
], [4.2, 7.0, 6.3])

# ----------------------------------------------------------------------------- 2
h("2. Ideación del producto de datos")
para("**Usuarios y dolores.** (i) *Coordinación de agenda/call center*: no sabe a quién llamar primero y sobreagenda a ciegas. (ii) *Dirección "
     "de MinHalma*: pierde capacidad facturable y no mide el efecto de las estrategias de recordatorio. (iii) *Psicólogos, nutricionistas y "
     "fisioterapeutas*: huecos en la agenda y pacientes que interrumpen el plan de sesiones. (iv) *Pacientes*: listas de espera más largas por cupos perdidos.")
para("**Componentes (4 de 4):** (1) **modelo de ML** de riesgo por cita, calibrado y explicable (SHAP); (2) **API REST** (FastAPI) con "
     "endpoints /score y /explain consumidos por el sistema de agendamiento; (3) **dashboard** (Streamlit) con lista priorizada, acción sugerida y "
     "KPI; (4) integración opcional con mensajería para recordatorios focalizados. **Requerimientos clave:** puntaje al agendar y en D-3; latencia "
     "< 500 ms por cita; seudonimización y control de acceso por rol; registro de la acción tomada y su desenlace (para medir impacto y reentrenar); "
     "monitoreo de calibración y deriva. **Restricción de uso:** el puntaje solo activa acciones de apoyo (confirmar, recordar, reasignar cupos "
     "liberados); nunca se usa para negar o postergar una cita.")
figure(ROOT / "mockup/mockup_dashboard.png", 15.5, "Figura 1. Mockup del tablero operativo de NeuroAgenda (datos ficticios). Fuente: elaboración propia.")

# ----------------------------------------------------------------------------- 3
h("3. Aspectos éticos y legales")
table([
    ["Riesgo", "Marco aplicable", "Mitigación en el producto"],
    ["Datos de salud = datos sensibles", "Ley 1581 de 2012 (arts. 5–6) y Decreto 1377 de 2013: autorización previa, expresa e informada y finalidad explícita",
     "Autorización en el aviso de privacidad de MinHalma; minimización (solo variables de agenda); seudonimización en origen"],
    ["Reserva de historia clínica y secreto profesional (crítico en salud mental)", "Resolución 1995 de 1999; Ley 1090 de 2006 (psicología); Ley 528 de 1999 (fisioterapia); Ley 73 de 1979 (nutrición); Ley 1616 de 2013",
     "El modelo no usa diagnósticos ni notas clínicas; el tablero no muestra nombre ni diagnóstico; acceso por rol y auditoría"],
    ["Interoperabilidad y custodia", "Ley 2015 de 2020 (historia clínica electrónica interoperable)", "Acuerdo de transferencia y confidencialidad con MinHalma; los datos permanecen en su infraestructura"],
    ["Sesgo y discriminación algorítmica", "CONPES 3975 de 2019 (IA confiable); evidencia de sesgo en algoritmos de salud (Obermeyer et al., 2019)",
     "Variables como subsidio y barrio son proxies socioeconómicos: auditoría de equidad por subgrupo y prohibición de uso restrictivo"],
    ["Investigación con datos secundarios", "Resolución 8430 de 1993 (investigación sin riesgo)", "Aval de un comité de ética (el de la universidad, si MinHalma no cuenta con uno) antes de usar datos reales"],
    ["Transferibilidad del conjunto proxy", "Datos públicos y anonimizados (Kaggle); contexto Brasil 2016", "Uso solo exploratorio; el modelo final se entrena con datos de MinHalma"],
], [4.3, 6.9, 6.3], size=9)

# ----------------------------------------------------------------------------- 4
h("4. Enfoque analítico")
para("**Pregunta de negocio:** ¿a qué citas debe dirigirse la confirmación activa para evitar el mayor número de inasistencias con una "
     "capacidad de contacto limitada? Se formula como **clasificación binaria supervisada** (no-show = 1) con desbalance moderado "
     f"({pct(R['tasa_no_show'])}). Hipótesis (estado preliminar según el EDA, sección 6):")
table([
    ["Hipótesis", "Técnica de contraste", "Evidencia preliminar"],
    ["H1. A mayor antelación entre agendamiento y cita, mayor inasistencia", "Mann-Whitney, χ², OR ajustado",
     f"Soportada: {pct(ta['0'])} (mismo día) vs. {pct(ta['31–60'])} (31–60 d)"],
    ["H2. El historial de inasistencias del paciente es predictor independiente", "Regresión logística, tasas por historial",
     f"Soportada: {orf('inasist_previas_ge1')}"],
    ["H3. El recordatorio SMS reduce la inasistencia una vez ajustado por antelación", "OR de Mantel-Haenszel estratificado",
     f"Soportada: OR crudo {num(R['sms_or_crudo'])} → OR MH {num(R['sms_or_mh'][0])}"],
    ["H4. El riesgo varía por edad y nivel socioeconómico (alerta de equidad)", "χ², V de Cramér, OR ajustados",
     f"Soportada: 12–17 años {pct(te['12–17'])} vs. 60–74 {pct(te['60–74'])}; subsidio {orf('subsidio')}"],
    ["H5. Las variables de agenda bastan para priorizar contactos mejor que el azar", "Validación cruzada agrupada por paciente",
     f"Soportada: ROC-AUC {num(best.roc_media, 3)}; contactar el 30 % de mayor riesgo captura {pct(best.recall_top30_media, 0)} de los no-show"],
], [6.3, 4.4, 6.8], size=9)
para("**Técnicas (Entrega 2):** regresión logística regularizada (línea base interpretable), Random Forest y gradient boosting "
     "(LightGBM/XGBoost), que la literatura reporta como los de mejor desempeño en este problema (Salazar et al., 2022), con búsqueda de hiperparámetros; calibración isotónica; explicabilidad con SHAP. **Validación:** GroupKFold por paciente "
     "(evita fuga entre citas de una misma persona) y partición temporal (últimas semanas como prueba). **Métricas del modelo:** PR-AUC "
     "(principal, por el desbalance), ROC-AUC, recall y precisión en el top-k % (capacidad real de llamadas), Brier y curva de calibración, y "
     "brecha de recall entre subgrupos. El umbral se elegirá por costo esperado (costo del contacto vs. valor del cupo recuperado).")

# ----------------------------------------------------------------------------- 5
h("5. Recolección de datos")
para(f"**Fuente 1 (disponible, proxy):** *Medical Appointment No Shows* (Hoppen, 2016), CSV estructurado de {ent(R['n_citas'])} citas × 14 atributos de "
     f"{ent(R['n_pacientes'])} pacientes del sistema público de Vitória (Brasil), citas entre {R['periodo_citas'][0]} y {R['periodo_citas'][1]}. Permite "
     "validar hipótesis y pipeline mientras se formaliza el acceso. **Fuente 2 (Entrega 2):** extracto seudonimizado del sistema de agendamiento de "
     "MinHalma (cita, servicio, profesional, hora, canal, estado asistió/canceló/no asistió, edad, sexo, régimen, municipio). **Fuente 3:** calendario de "
     "festivos de Colombia. Diccionario completo en el repositorio (docs/diccionario_datos.md); atributos principales:")
table([
    ["Atributo (original)", "Tipo", "Descripción", "Relevancia"],
    ["AppointmentID / PatientId", "ID", "Identificador de cita / paciente (seudónimo)", "Agrupación y historial"],
    ["ScheduledDay / AppointmentDay", "Fecha", "Fecha-hora de agendamiento / fecha de la cita (sin hora)", "Deriva la antelación (predictor principal)"],
    ["Age, Gender", "Num. / Cat.", "Edad en años; sexo (F/M)", "Riesgo y auditoría de equidad"],
    ["Scholarship", "Binaria", "Beneficiario de Bolsa Família (proxy socioeconómico)", "Riesgo y equidad"],
    ["Hipertension, Diabetes, Alcoholism, Handcap", "Binaria / 0–4", "Condiciones registradas; Handcap = n.º de discapacidades", "Comorbilidad (sensible)"],
    ["SMS_received", "Binaria", "Se envió al menos un SMS de recordatorio", "Intervención evaluable"],
    ["Neighbourhood", "Cat. (81)", "Barrio de la unidad de atención", "Contexto territorial"],
    ["No-show", "Binaria", "Yes = no asistió (variable objetivo)", "Objetivo"],
], [5.0, 2.4, 6.1, 4.0], size=9)

# ----------------------------------------------------------------------------- 6
q = R["calidad"]
h("6. Entendimiento y calidad de los datos")
para(f"**Calidad.** Completitud total (0 nulos) y unicidad de cita_id. Se detectaron: {q['fecha_cita anterior a fecha_agendamiento']} citas con fecha "
     f"anterior al agendamiento y {q['edad < 0']} edad negativa (excluidas, {R['registros_excluidos']} registros); {q['edad > 100']} edades > 100 (marcadas para "
     f"validar); Handcap codificado 0–4 y no binario ({q['discapacidad fuera de {0,1} (codificada 0–4)']} casos); {ent(q['Filas duplicadas sin cita_id'])} registros "
     "idénticos salvo el ID (posible doble agendamiento, a revisar con el stakeholder); y la fecha de la cita **no tiene hora**, lo que impide modelar la "
     f"franja horaria. **Univariado:** la antelación es muy asimétrica (mediana {num(R['antelacion']['mediana'], 0)} días, asimetría {num(R['antelacion']['asimetria'])}; "
     f"{pct(R['antelacion']['pct_mismo_dia'])} se agenda el mismo día); {pct(R['citas_por_paciente']['pct_pacientes_1_cita'])} de los pacientes tiene una sola cita "
     f"(máximo {R['citas_por_paciente']['max']}), y {pct(R['pct_sms'])} recibió SMS.")
figure(ROOT / "reports/figures/fig2_bivariado.png", 17.3, "Figura 2. Tasa de inasistencia (IC 95 % de Wilson) por antelación, edad, día y SMS. Línea punteada = tasa global.")
b = R["bivariado"]
para(f"**Bivariado.** La antelación es la variable con mayor tamaño de efecto (r rank-biserial = {num(b['antelacion_dias'])}): la inasistencia pasa de "
     f"{pct(ta['0'])} en citas del mismo día a {pct(ta['8–14'])}–{pct(ta['31–60'])} cuando se agenda con 8 a 60 días. La edad muestra un gradiente "
     f"inverso (máximo en 12–17 años, {pct(te['12–17'])}). Sexo (p = 0,17) y día de la semana tienen efectos despreciables (V < 0,02).")
para(f"**Multivariado.** El SMS parece *aumentar* la inasistencia (OR crudo = {num(R['sms_or_crudo'])}), pero solo se envía a citas con ≥ "
     f"{R['sms_min_antelacion']} días de antelación, que de por sí tienen más riesgo: al estratificar por días de antelación el efecto se invierte "
     f"(**OR de Mantel-Haenszel = {num(R['sms_or_mh'][0])} [{num(R['sms_or_mh'][1])}–{num(R['sms_or_mh'][2])}]**), una paradoja de Simpson que justifica "
     f"focalizar recordatorios. En el modelo logístico ajustado, el historial de inasistencias ({orf('inasist_previas_ge1')}), el subsidio, la edad y la "
     f"antelación son los efectos independientes más robustos; la tasa sube de {pct(th['0'])} sin inasistencias previas a {pct(th['3+'])} con 3 o más. "
     f"{R['barrios_fuera_998']} de {R['n_barrios']} barrios se salen de los límites del 99,8 % del funnel plot (variación territorial real).")
figure(ROOT / "reports/figures/fig3_multivariado.png", 17.3, "Figura 3. Correlaciones de Spearman, OR ajustados (IC 95 %) y tasa según inasistencias previas sin fuga temporal.")
para(f"**Suficiencia.** Una línea base sin optimizar (validación cruzada agrupada por paciente, k = 5) alcanza ROC-AUC = {num(lr.roc_media, 3)} con "
     f"regresión logística y {num(best.roc_media, 3)} ± {num(best.roc_de, 3)} con gradient boosting e historial; PR-AUC = {num(best.pr_media, 3)} "
     f"({num(best.pr_media / R['tasa_no_show'], 1)} veces la prevalencia) y Brier = {num(best.brier_media, 3)} (vs. {num(R['tasa_no_show'] * (1 - R['tasa_no_show']), 3)} "
     "de predecir la prevalencia). Al construir el historial se corrigió una fuga temporal (contar citas agendadas antes pero aún no ocurridas), que "
     "inflaba el ROC-AUC en ≈ 0,01. La señal es **suficiente para priorizar, pero moderada**: faltan servicio, hora, profesional, distancia y "
     "cancelaciones previas, que se solicitarán a MinHalma. Detalle: notebooks/01_eda_calidad.ipynb.")

# ----------------------------------------------------------------------------- 7
h("7. Conclusiones iniciales y próximos pasos")
bullet(f"**La antelación es la palanca operativa principal:** confirmar activamente en D-3/D-1 las citas agendadas con 8 a 60 días de antelación, donde la inasistencia supera el {pct(min(ta['8–14'], ta['15–30'], ta['31–60']), 0)}.")
bullet(f"**Los recordatorios funcionan si se mide bien:** ajustado por antelación, el SMS reduce las odds de inasistencia ≈ {pct(1 - R['sms_or_mh'][0], 0)}; la comparación cruda lleva a la conclusión opuesta.")
bullet(f"**Focalizar es viable:** contactar el 30 % de citas de mayor riesgo capturaría ≈ {pct(best.recall_top30_media, 0)} de las inasistencias, frente al 30 % esperado al azar.")
bullet("**Equidad como requisito de diseño:** jóvenes y beneficiarios de subsidio presentan más riesgo; el modelo debe ampliar el apoyo a estos grupos, no restringir su acceso.")
bullet("**Acciones siguientes:** formalizar el acuerdo de datos y el aval ético con MinHalma (interacción 1 de la bitácora); replicar el EDA con datos reales; "
       "en la Entrega 2, preparar datos, entrenar y comparar ≥ 3 algoritmos, y desplegar API y dashboard.")

# ----------------------------------------------------------------------------- referencias
h("Referencias")
refs = [
    "Blu Radio. (2025, 15 de julio). Preocupación por inasistencia a más de 15.000 citas médicas en hospital de San Gil. https://www.bluradio.com/regiones/santanderes/preocupacion-por-inasistencia-a-mas-de-15-000-citas-medicas-en-hospital-de-san-gil-rg10",
    "Congreso de Colombia. (1979). Ley 73 de 1979, por la cual se reglamenta la profesión de nutricionista-dietista.",
    "Congreso de Colombia. (1999). Ley 528 de 1999, por la cual se reglamenta el ejercicio de la profesión de fisioterapia.",
    "Congreso de Colombia. (2006). Ley 1090 de 2006, por la cual se reglamenta el ejercicio de la profesión de Psicología.",
    "Congreso de Colombia. (2012). Ley Estatutaria 1581 de 2012, por la cual se dictan disposiciones generales para la protección de datos personales.",
    "Congreso de Colombia. (2013). Ley 1616 de 2013, por medio de la cual se expide la ley de salud mental.",
    "Congreso de Colombia. (2020). Ley 2015 de 2020, por medio de la cual se crea la historia clínica electrónica interoperable.",
    "Consejo Nacional de Política Económica y Social. (2019). Documento CONPES 3975: Política nacional para la transformación digital e inteligencia artificial. DNP.",
    "Dantas, L. F., Fleck, J. L., Cyrino Oliveira, F. L., & Hamacher, S. (2018). No-shows in appointment scheduling – a systematic literature review. Health Policy, 122(4), 412–421. https://doi.org/10.1016/j.healthpol.2018.02.002",
    "Hoppen, J. (2016). Medical Appointment No Shows [Conjunto de datos]. Kaggle. https://www.kaggle.com/datasets/joniarroba/noshowappointments",
    "Ministerio de Salud. (1993). Resolución 8430 de 1993, normas científicas, técnicas y administrativas para la investigación en salud.",
    "Ministerio de Salud. (1999). Resolución 1995 de 1999, normas para el manejo de la historia clínica.",
    "Obermeyer, Z., Powers, B., Vogeli, C., & Mullainathan, S. (2019). Dissecting racial bias in an algorithm used to manage the health of populations. Science, 366(6464), 447–453. https://doi.org/10.1126/science.aax2342",
    "Presidencia de la República de Colombia. (2013). Decreto 1377 de 2013, por el cual se reglamenta parcialmente la Ley 1581 de 2012.",
    "Salazar, L. H. A., Leithardt, V. R. Q., Parreira, W. D., da Rocha Fernandes, A. M., Barbosa, J. L. V., & Correia, S. D. (2022). Application of machine learning techniques to predict a patient's no-show in the healthcare sector. Future Internet, 14(1), 3. https://doi.org/10.3390/fi14010003",
    "Valora Analitik. (2026, 20 de abril). Calidad de la salud empeoró en 2025 y los colombianos asisten cada vez menos a citas médicas. https://www.valoraanalitik.com/calidad-de-la-salud-empeoro-en-2025-y-los-colombianos-asisten-cada-vez-menos-a-citas-medicas/",
]
for ref in refs:
    p = para(ref, size=9, align=WD_ALIGN_PARAGRAPH.LEFT, after=1)
    p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)

out = ROOT / "reports" / "Entrega1_Reporte.docx"
doc.save(out); print("Reporte:", out)
